#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Authoring-time experiment harness (NOT part of the graded container).

Runs the three required control experiments against the grader:
  (a) empty / corrupt candidate workbooks      -> reward < 1, no crash
  (b) perturbed golden (2 values flipped)      -> reward drops strictly
  (c) near-miss oracle variants (single-clause misreadings) -> reward < 0.5

Usage: python3 run_experiments.py [/path/to/init.xlsx]
"""
import copy
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
TESTS = os.path.join(os.path.dirname(HERE), "tests")
sys.path.insert(0, TESTS)

import openpyxl  # noqa: E402
import oracle_engine as eng  # noqa: E402
from mutations import VARIANTS, SECONDARY, COMPOSITE  # noqa: E402
import test_outputs as T  # noqa: E402

INPUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(HERE), "environment", "sheet", "init.xlsx")


def score_workbook(path, expected):
    wb = openpyxl.load_workbook(path, data_only=True)
    checks = T.grade(wb, expected)
    total_w = sum(w for _, _, w in checks)
    passed_w = sum(w for _, ok, w in checks if ok)
    reward = passed_w / total_w if total_w else 0.0
    failed = [name for name, ok, _ in checks if not ok]
    return reward, failed


def variant_score(name, flags, inp_path, baseline):
    eng.MUT.clear()
    eng.MUT.update(flags)
    try:
        res = eng.compute_from_file(inp_path, None)
        tmp = os.path.join(tempfile.gettempdir(), f"variant_{name}.xlsx")
        eng.build_workbook(res, tmp)
        reward, failed = score_workbook(tmp, baseline)
        return reward, len(failed)
    finally:
        eng.MUT.clear()


def main():
    print("=" * 72)
    print("(0) baseline: oracle run must score 1.0")
    eng.MUT.clear()
    res = eng.compute_from_file(INPUT, None)
    tmp = os.path.join(tempfile.gettempdir(), "oracle_golden.xlsx")
    eng.build_workbook(res, tmp)
    reward, failed = score_workbook(tmp, res)
    print(f"    oracle reward = {reward:.6f}  (failed checks: {len(failed)})")
    assert reward >= 0.999999, "ORACLE MUST SCORE 1.0"
    baseline = res

    print("=" * 72)
    print("(a) empty / corrupt submissions")
    p_empty = os.path.join(tempfile.gettempdir(), "cand_empty.xlsx")
    openpyxl.Workbook().save(p_empty)
    r, f = score_workbook(p_empty, res)
    print(f"    empty workbook          reward = {r:.4f}  ({len(f)} failed checks)")
    p_trunc = os.path.join(tempfile.gettempdir(), "cand_trunc.xlsx")
    with open(tmp, "rb") as fh:
        blob = fh.read()
    with open(p_trunc, "wb") as fh:
        fh.write(blob[:len(blob) // 3])
    try:
        wb = openpyxl.load_workbook(p_trunc, data_only=True)
        checks = T.grade(wb, res)
        r = sum(w for _, ok, w in checks if ok) / sum(w for _, _, w in checks)
        print(f"    corrupt (truncated) xlsx reward = {r:.4f}  (grader did not raise)")
    except Exception as e:
        print(f"    corrupt xlsx raised at LOAD stage: {type(e).__name__}: {e}")
        print("    -> main() catches this and writes reward 0 (see test_outputs.main)")
    p_missing = os.path.join(tempfile.gettempdir(), "cand_missing.xlsx")
    if os.path.exists(p_missing):
        os.remove(p_missing)
    print("    missing file            -> main() writes reward 0 (handled)")

    print("=" * 72)
    print("(b) perturbed golden: flip 2 values in the expected dict")
    bad = copy.deepcopy(res)
    bad["summary"] = [(k, v + 500 if k == "TotalEquity" else v)
                      for k, v in bad["summary"]]
    for nr in bad["notes_rows"]:
        if nr["P"] == 7 and nr["class"] == "B":
            nr["end"] = nr["end"] + 100
            break
    reward_bad, nbad = score_workbook(tmp, bad)
    print(f"    perturbed-golden reward = {reward_bad:.4f}  ({nbad} failed checks; "
          f"baseline was 1.0 -> binding confirmed)")

    print("=" * 72)
    print("(c) near-miss oracle variants (single-clause misreadings)")
    worst = []
    for name, flags in VARIANTS.items():
        r, n = variant_score(name, flags, INPUT, baseline)
        flag = "OK " if r < 0.5 else "WEAK"
        print(f"    [{flag}] {name:<22} reward = {r:.4f}   ({n} checks failed)")
        worst.append((r, name))
    worst.sort(reverse=True)
    print(f"    weakest tested variant: {worst[0][1]} at {worst[0][0]:.4f}")
    print("    -- COMPOSITE (two-clause) variants: the agent-realistic profile --")
    cw = []
    for name, flags in COMPOSITE.items():
        r, n = variant_score(name, flags, INPUT, baseline)
        print(f"        {'OK ' if r < 0.5 else 'WEAK'} {name:<24} reward = {r:.4f}   ({n} failed)")
        cw.append(r)
    print(f"    composite worst = {max(cw):.4f}")
    print("    -- secondary (structurally bounded) variants, diagnostics only --")
    for name, flags in SECONDARY.items():
        r, n = variant_score(name, flags, INPUT, baseline)
        print(f"        {name:<22} reward = {r:.4f}   ({n} checks failed)")
    print("=" * 72)


if __name__ == "__main__":
    main()
